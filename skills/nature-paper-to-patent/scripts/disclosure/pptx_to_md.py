#!/usr/bin/env python3
r"""Export PowerPoint (.pptx) to Markdown slide by slide, extracting embedded images
for Step 2 scanning and agent reading. Requires python-pptx (requirements.txt).
Usage: python pptx_to_md.py -i input.pptx -o output.md
Images default to <markdown-stem>_media/ beside the output Markdown file.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        print(
            "Missing python-pptx; install skill requirements in the selected environment.",
            file=sys.stderr,
        )
        sys.exit(1)
    return Presentation, MSO_SHAPE_TYPE


def _walk_shapes(shapes, MSO_SHAPE_TYPE):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes, MSO_SHAPE_TYPE)
        else:
            yield shape


def _shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        t = (shape.text_frame.text or "").strip()
        return t
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = []
            for cell in row.cells:
                cells.append((cell.text or "").strip().replace("\n", " "))
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            return "\n".join(rows)
    return ""


def _rel_media_path(out_file: Path, media_file: Path) -> str:
    try:
        return media_file.relative_to(out_file.parent).as_posix()
    except ValueError:
        return media_file.as_posix()


def _run(input_pptx: Path, output_md: Path, media_dir: Path | None) -> int:
    Presentation, MSO_SHAPE_TYPE = _require_pptx()

    if not input_pptx.is_file():
        print(f"Input file does not exist: {input_pptx}", file=sys.stderr)
        return 2
    suf = input_pptx.suffix.lower()
    if suf not in (".pptx", ".ppsx"):
        print("Warning: expected .pptx/.ppsx (OOXML); legacy .ppt is unsupported.", file=sys.stderr)

    output_md = output_md.resolve()
    output_md.parent.mkdir(parents=True, exist_ok=True)

    if media_dir is None:
        media_dir = output_md.parent / f"{output_md.stem}_media"
    else:
        media_dir = media_dir.resolve()
    media_dir.mkdir(parents=True, exist_ok=True)

    try:
        prs = Presentation(str(input_pptx))
    except Exception as e:
        print(f"Could not open presentation: {e}", file=sys.stderr)
        return 3

    lines: list[str] = [
        f"<!-- Converted from {input_pptx.name} by pptx_to_md.py; do not edit this metadata line -->\n"
    ]
    img_counter = [0]

    for sn, slide in enumerate(prs.slides, start=1):
        lines.append(f"\n## Slide {sn}\n")

        for shape in _walk_shapes(slide.shapes, MSO_SHAPE_TYPE):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = shape.image
                    ext = (img.ext or "png").lower()
                    if ext == "jpeg":
                        ext = "jpg"
                    img_counter[0] += 1
                    fname = f"slide{sn:02d}_img{img_counter[0]:04d}.{ext}"
                    out_img = media_dir / fname
                    out_img.write_bytes(img.blob)
                    rel = _rel_media_path(output_md, out_img)
                    lines.append(f"\n![]({rel})\n")
                except Exception as e:
                    print(f"Warning: image extraction failed on slide {sn}: {e}", file=sys.stderr)
                continue

            block = _shape_text(shape)
            if block:
                lines.append(block)
                lines.append("\n\n")

        try:
            nf = slide.notes_slide.notes_text_frame
            note_txt = (nf.text or "").strip() if nf is not None else ""
            if note_txt:
                lines.append("\n**Notes**:\n\n")
                lines.append(note_txt)
                lines.append("\n\n")
        except (AttributeError, ValueError):
            pass

    body = "".join(lines).rstrip() + "\n"
    output_md.write_text(body, encoding="utf-8")

    print(f"Written: {output_md}")
    print(f"Image directory: {media_dir}")
    return 0


def main() -> int:
    p = argparse.ArgumentParser(description="PowerPoint (.pptx) to Markdown with extracted images")
    p.add_argument("-i", "--input", required=True, type=Path, help="Input .pptx/.ppsx path")
    p.add_argument("-o", "--output", required=True, type=Path, help="Output .md path")
    p.add_argument(
        "--media-dir",
        type=Path,
        default=None,
        help="Image directory (default: <markdown-stem>_media beside the .md file)",
    )
    args = p.parse_args()
    return _run(args.input, args.output, args.media_dir)


if __name__ == "__main__":
    raise SystemExit(main())
